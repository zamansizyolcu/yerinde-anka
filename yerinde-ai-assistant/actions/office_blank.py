"""
actions/office_blank.py — 'Word aç' dendiğinde başlangıç ekranı yerine
DOĞRUDAN boş bir belge açar: masaüstündeki 'Çalışmalarım' klasörüne boş
.docx/.xlsx/.pptx oluşturup varsayılan Office uygulamasıyla açar
(MS Office / LibreOffice / OnlyOffice hangisi kuruluysa o açılır).
"""

from __future__ import annotations

import platform
import subprocess
import time
from pathlib import Path

from actions.code_tools import ensure_workspace_folder

_IS_WINDOWS = platform.system() == "Windows"

# Türkçe takma adlar: "sunum aç" da boş sunum açsın (eskiden 'sunum' diye bir
# uygulama aranıyor ve PowerPoint başlangıç ekranı açılıyordu).
OFFICE_ALIASES = {
    "sunum": "powerpoint", "sunumu": "powerpoint", "slayt": "powerpoint",
    "powerpoint": "powerpoint", "power point": "powerpoint",
    "belge": "word", "belgeyi": "word", "yazı": "word", "word": "word",
    "tablo": "excel", "hesap tablosu": "excel", "excel": "excel",
}


def normalize_kind(name: str) -> str:
    """'sunum' → 'powerpoint' gibi çevirir; tanımadığını olduğu gibi döner."""
    return OFFICE_ALIASES.get((name or "").lower().strip(), (name or "").lower().strip())


KINDS = {
    "word": (".docx", "Yeni Belge"),
    "excel": (".xlsx", "Yeni Tablo"),
    "powerpoint": (".pptx", "Yeni Sunum"),
}


def _make_blank(kind: str, path: Path) -> bool:
    try:
        if kind == "word":
            from docx import Document
            Document().save(str(path))
        elif kind == "excel":
            from openpyxl import Workbook
            Workbook().save(str(path))
        else:
            from pptx import Presentation
            Presentation().save(str(path))
        return True
    except Exception:
        return False


def _open_file(path: Path) -> None:
    if _IS_WINDOWS:
        import os
        os.startfile(str(path))  # noqa: S606
    else:
        subprocess.Popen(["xdg-open", str(path)],
                         stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)


def create_blank_document(kind: str) -> str | None:
    """Başarıda kullanıcıya okunacak mesaj, desteklenmeyen/kütüphanesizse None
    (çağıran open_app'e düşer)."""
    kind = (kind or "").lower().strip()
    if kind not in KINDS:
        return None
    ext, base = KINDS[kind]
    folder = ensure_workspace_folder()
    stamp = time.strftime("%Y-%m-%d %H.%M")
    path = folder / f"{base} {stamp}{ext}"
    i = 2
    while path.exists():
        path = folder / f"{base} {stamp} ({i}){ext}"
        i += 1
    if not _make_blank(kind, path):
        return None
    try:
        _open_file(path)
    except Exception as e:
        return f"Boş belge oluşturuldu ({path}) ama açılamadı: {e}"
    return f"Boş {base.lower()} açıldı: {path.name} (Çalışmalarım klasörüne kaydedildi)."
